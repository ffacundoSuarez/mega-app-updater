"""
Compara dos informes .pptx de Brand Audit gráfico por gráfico.

Sirve como test de regresión del motor: el motor reemplaza la columna de la ola
si ya existe en el gráfico (en vez de agregar una nueva), así que correrlo sobre
el informe que generó Chris, con la misma ola y la misma base, tiene que
reproducir sus números.

Compara, para cada shape con gráfico o tabla:
  - que exista en ambos archivos
  - las categorías del eje
  - los valores de cada serie, con tolerancia configurable
  - el contenido de las celdas de las tablas

Uso:
    python scripts/compare-ppt.py <referencia.pptx> <generado.pptx> [--tol 0.01]
    python scripts/compare-ppt.py ref.pptx gen.pptx --solo-ola "Ago 26"
"""

import argparse
import sys
from collections import OrderedDict

from pptx import Presentation


def limpiar(texto):
    """Normaliza etiquetas: saca zero-width spaces y espacios de sobra."""
    return str(texto).replace("\u200b", "").strip()


def recolectar(path):
    """Devuelve {nombre_shape: {"tipo", "categorias", "series"|"celdas"}}."""
    prs = Presentation(path)
    encontrados = OrderedDict()

    def visitar(shapes):
        for sh in shapes:
            try:
                nombre = sh.name
            except Exception:
                continue

            # Grupos: bajamos recursivamente (shape_type 6 == GROUP)
            if sh.shape_type == 6:
                try:
                    visitar(sh.shapes)
                except Exception:
                    pass
                continue

            if getattr(sh, "has_chart", False):
                try:
                    chart = sh.chart
                    cats = [limpiar(c) for c in chart.plots[0].categories]
                    series = OrderedDict()
                    for s in chart.series:
                        clave = limpiar(s.name) if s.name else f"serie_{len(series)}"
                        series[clave] = list(s.values)
                    encontrados[nombre] = {
                        "tipo": "chart",
                        "categorias": cats,
                        "series": series,
                    }
                except Exception as e:
                    encontrados[nombre] = {"tipo": "chart", "error": str(e)}

            elif getattr(sh, "has_table", False):
                try:
                    celdas = [
                        [limpiar(c.text) for c in fila.cells] for fila in sh.table.rows
                    ]
                    encontrados[nombre] = {"tipo": "tabla", "celdas": celdas}
                except Exception as e:
                    encontrados[nombre] = {"tipo": "tabla", "error": str(e)}

    for slide in prs.slides:
        visitar(slide.shapes)
    return encontrados


def casi_igual(a, b, tol):
    """True si dos valores numéricos (o ambos vacíos) coinciden dentro de tol."""
    if a is None and b is None:
        return True
    if a is None or b is None:
        return False
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return a == b


def comparar_chart(ref, gen, tol, solo_ola):
    """Devuelve lista de diferencias legibles para un gráfico."""
    difs = []

    if ref.get("categorias") != gen.get("categorias"):
        difs.append(
            "    categorías distintas\n"
            f"      ref: {ref.get('categorias')}\n"
            f"      gen: {gen.get('categorias')}"
        )
        # Sin ejes alineados no tiene sentido comparar valores posicionalmente.
        return difs

    cats = ref.get("categorias", [])
    # Si sólo interesa una ola, comparamos únicamente esa posición del eje.
    indices = list(range(len(cats)))
    if solo_ola:
        objetivo = limpiar(solo_ola)
        indices = [i for i, c in enumerate(cats) if c == objetivo]
        if not indices:
            return [f"    la ola {solo_ola!r} no está en el eje: {cats}"]

    for serie, vals_ref in ref.get("series", {}).items():
        if serie not in gen.get("series", {}):
            difs.append(f"    falta la serie {serie!r} en el generado")
            continue
        vals_gen = gen["series"][serie]
        for i in indices:
            vr = vals_ref[i] if i < len(vals_ref) else None
            vg = vals_gen[i] if i < len(vals_gen) else None
            if not casi_igual(vr, vg, tol):
                etiqueta = cats[i] if i < len(cats) else f"#{i}"
                difs.append(f"    {serie!r} @ {etiqueta!r}: ref={vr} vs gen={vg}")

    extras = set(gen.get("series", {})) - set(ref.get("series", {}))
    for s in sorted(extras):
        difs.append(f"    serie de más en el generado: {s!r}")

    return difs


def comparar_tabla(ref, gen):
    """Devuelve lista de diferencias legibles para una tabla."""
    difs = []
    fr, fg = ref.get("celdas", []), gen.get("celdas", [])
    if len(fr) != len(fg):
        difs.append(f"    cantidad de filas: ref={len(fr)} vs gen={len(fg)}")
    for i, (fila_r, fila_g) in enumerate(zip(fr, fg)):
        if fila_r != fila_g:
            for j, (cr, cg) in enumerate(zip(fila_r, fila_g)):
                if cr != cg:
                    difs.append(f"    celda [{i}][{j}]: ref={cr!r} vs gen={cg!r}")
    return difs


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("referencia", help="informe de referencia (ej: el de Chris)")
    ap.add_argument("generado", help="informe generado por la app")
    ap.add_argument(
        "--tol",
        type=float,
        default=0.005,
        help="tolerancia numérica absoluta (default 0.005)",
    )
    ap.add_argument(
        "--solo-ola",
        default=None,
        help="comparar sólo esta categoría del eje (ej: 'Ago 26')",
    )
    ap.add_argument(
        "--max",
        type=int,
        default=8,
        help="máximo de diferencias a mostrar por shape (default 8)",
    )
    args = ap.parse_args()

    print(f"referencia: {args.referencia}")
    print(f"generado:   {args.generado}")
    extra = f" | sólo ola {args.solo_ola!r}" if args.solo_ola else ""
    print(f"tolerancia: {args.tol}{extra}")
    print("=" * 78)

    ref = recolectar(args.referencia)
    gen = recolectar(args.generado)

    solo_ref = [n for n in ref if n not in gen]
    solo_gen = [n for n in gen if n not in ref]

    con_difs = 0
    iguales = 0

    for nombre, datos_ref in ref.items():
        if nombre not in gen:
            continue
        datos_gen = gen[nombre]

        if datos_ref.get("tipo") != datos_gen.get("tipo"):
            print(f"\n[TIPO] {nombre}: ref={datos_ref.get('tipo')} gen={datos_gen.get('tipo')}")
            con_difs += 1
            continue

        if datos_ref.get("tipo") == "chart":
            difs = comparar_chart(datos_ref, datos_gen, args.tol, args.solo_ola)
        else:
            difs = comparar_tabla(datos_ref, datos_gen)

        if difs:
            con_difs += 1
            print(f"\n[DIF] {nombre}  ({len(difs)} diferencias)")
            for d in difs[: args.max]:
                print(d)
            if len(difs) > args.max:
                print(f"    ... y {len(difs) - args.max} más")
        else:
            iguales += 1

    print("\n" + "=" * 78)
    print(f"shapes comparados : {iguales + con_difs}")
    print(f"  identicos       : {iguales}")
    print(f"  con diferencias : {con_difs}")
    if solo_ref:
        print(f"sólo en referencia ({len(solo_ref)}): {', '.join(solo_ref[:15])}")
    if solo_gen:
        print(f"sólo en generado   ({len(solo_gen)}): {', '.join(solo_gen[:15])}")

    return 1 if (con_difs or solo_ref or solo_gen) else 0


if __name__ == "__main__":
    sys.exit(main())
